import streamlit as st
from urllib.parse import quote
from openpyxl import load_workbook

RUTA = "Fixture_Valorant_uLeague_Clausura_2026.xlsx"
HOJA = "Equipos"


def leer_equipos(ruta):
    ws = load_workbook(ruta, read_only=True, data_only=True)[HOJA]
    filas = [list(f) for f in ws.iter_rows(values_only=True)]

    equipos = {}

    for r, fila in enumerate(filas):
        for c, celda in enumerate(fila):
            if not celda or "#" in str(celda):
                continue

            arriba = (
                filas[r - 1][c]
                if r > 0 and c < len(filas[r - 1])
                else None
            )

            abajo = (
                filas[r + 1][c]
                if r + 1 < len(filas) and c < len(filas[r + 1])
                else None
            )

            if (arriba and str(arriba).strip()) or not abajo or "#" not in str(abajo):
                continue

            jugadores = []
            i = r + 1

            while i < len(filas) and c < len(filas[i]) and filas[i][c]:
                jugadores.append(str(filas[i][c]).strip())
                i += 1

            equipos[str(celda).strip()] = jugadores

    return equipos


def url_perfil(riot_id):
    nombre, tag = riot_id.split("#", 1)

    id_codificado = quote(
        f"{nombre.strip()}#{tag.replace(' ', '')}",
        safe=""
    )

    return (
        f"https://tracker.gg/valorant/profile/riot/"
        f"{id_codificado}/overview"
    )


# -----------------------------
# INTERFAZ
# -----------------------------

st.set_page_config(
    page_title="Valorant uLeague Tracker",
    page_icon="🎯",
    layout="centered"
)

st.title("🎯 Valorant uLeague Tracker")
st.write("Selecciona un equipo para ver sus perfiles de Tracker.gg.")

try:
    equipos = leer_equipos(RUTA)

except Exception as e:
    st.error(f"No se pudo leer el Excel: {e}")
    st.stop()


st.success(f"Se cargaron {len(equipos)} equipos.")

nombre_equipo = st.selectbox(
    "Selecciona un equipo",
    sorted(equipos.keys())
)

jugadores = equipos[nombre_equipo]

st.subheader(nombre_equipo)
st.write(f"**{len(jugadores)} jugadores**")


# Mostrar jugadores
urls = []

for riot_id in jugadores:

    if "#" not in riot_id:
        st.warning(f"⚠️ {riot_id} no tiene tag (#)")
        continue

    url = url_perfil(riot_id)
    urls.append(url)

    st.markdown(
        f"**{riot_id}**  \n"
        f"[🔗 Abrir perfil en Tracker.gg]({url})"
    )


st.divider()

if urls:
    st.info(
        "Puedes abrir cada perfil individualmente. "
        "El navegador puede bloquear la apertura automática de múltiples pestañas."
    )